import os
import sys
import json
import time
import threading
from collections import Counter
from datetime import datetime
from pathlib import Path
import logging
import hashlib
import shutil
import tempfile
from typing import Tuple, Optional, List

import customtkinter as ctk
from tkinter import filedialog, messagebox

from openpyxl import load_workbook, Workbook
from docx import Document
from pypdf import PdfReader, PdfWriter
from pptx import Presentation
from zipfile import ZipFile

# Carpetas/archivos que el propio script genera y que NUNCA deben volver a
# escanearse (esta era la causa de la duplicación de respaldos).
BACKUPS_DIRNAME = ".ficohsa_backups"
EXCLUDED_DIR_NAMES = {BACKUPS_DIRNAME.lower(), ".ficohsa_tagger", "$recycle.bin", "system volume information"}

# Atributos de archivo de Windows para detectar placeholders de OneDrive
# (archivos "solo en la nube" que no están descargados localmente).
FILE_ATTRIBUTE_OFFLINE = 0x00001000
FILE_ATTRIBUTE_RECALL_ON_DATA_ACCESS = 0x00400000

SETTINGS_FILE = "settings.json"
DEFAULT_BACKUP_RETENTION_DAYS = 30


def get_backups_base_root() -> Path:
    """Carpeta fija y visible para respaldos, fuera de OneDrive/SharePoint y
    fuera de AppData\\Local. AppData\\Local se descartó como ubicación: en
    pruebas reales una carpeta de respaldos ahí (con 345 archivos, 259MB)
    desapareció por completo sin dejar rastro en Defender, la Papelera de
    Reciclaje ni Acceso Controlado a Carpetas -- probablemente algún
    mecanismo de limpieza de Windows (Storage Sense) o de terceros que trata
    ese directorio como caché desechable. Se usa una carpeta visible en la
    raíz del perfil del usuario para que sea fácil de ubicar, verificar y
    excluir manualmente de cualquier limpiador de disco."""
    return Path.home() / "FicohsaBackups"


def format_size(num_bytes: float) -> str:
    for unit in ("B", "KB", "MB", "GB", "TB"):
        if num_bytes < 1024:
            return f"{num_bytes:.1f}{unit}"
        num_bytes /= 1024
    return f"{num_bytes:.1f}PB"


# ==================== CONFIGURACIÓN DE LOGGING ====================
class SecurityLogger:
    def __init__(self, log_dir: Path):
        self.log_dir = log_dir
        self.log_dir.mkdir(exist_ok=True)
        self.log_file = self.log_dir / f"etiquetado_seguridad_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"

        self.logger = logging.getLogger('FICOHSA_TAGGER')
        self.logger.setLevel(logging.DEBUG)

        fh = logging.FileHandler(self.log_file, encoding='utf-8')
        fh.setLevel(logging.DEBUG)

        formatter = logging.Formatter(
            '%(asctime)s | %(levelname)-8s | %(funcName)s | %(message)s',
            datefmt='%Y-%m-%d %H:%M:%S'
        )
        fh.setFormatter(formatter)
        self.logger.addHandler(fh)

    def info(self, msg): self.logger.info(msg)
    def error(self, msg): self.logger.error(msg)
    def warning(self, msg): self.logger.warning(msg)
    def debug(self, msg): self.logger.debug(msg)


# ==================== VALIDACIONES ====================
class SecurityValidator:
    MAX_FILE_SIZE = 100 * 1024 * 1024  # 100MB
    # Solo formatos que las librerías usadas pueden leer/escribir de verdad.
    # .xls/.doc/.ppt (binarios antiguos) se excluyen a propósito: openpyxl,
    # python-docx y python-pptx no pueden abrirlos y solo generaban errores.
    ALLOWED_EXTENSIONS = {'.xlsx', '.docx', '.pdf', '.pptx'}

    @staticmethod
    def validate_file(file_path: Path) -> Tuple[bool, str]:
        if not file_path.exists():
            return False, "Archivo no existe"

        if file_path.suffix.lower() not in SecurityValidator.ALLOWED_EXTENSIONS:
            return False, f"Extensión no permitida: {file_path.suffix}"

        try:
            file_size = file_path.stat().st_size
            if file_size > SecurityValidator.MAX_FILE_SIZE:
                max_mb = SecurityValidator.MAX_FILE_SIZE / 1024 / 1024
                return False, f"Archivo muy grande: {file_size / 1024 / 1024:.1f}MB (máx: {max_mb:.0f}MB)"

            if file_size == 0:
                return False, "Archivo vacío"
        except Exception as e:
            return False, f"Error al validar tamaño: {str(e)}"

        return True, "OK"

    @staticmethod
    def is_cloud_placeholder(file_path: Path) -> bool:
        """Detecta archivos de OneDrive/SharePoint que aún no se han descargado."""
        if sys.platform != 'win32':
            return False
        try:
            attrs = file_path.stat().st_file_attributes
            return bool(attrs & (FILE_ATTRIBUTE_OFFLINE | FILE_ATTRIBUTE_RECALL_ON_DATA_ACCESS))
        except (AttributeError, OSError):
            return False

    @staticmethod
    def calculate_hash(file_path: Path) -> str:
        sha256_hash = hashlib.sha256()
        with open(file_path, "rb") as f:
            for byte_block in iter(lambda: f.read(65536), b""):
                sha256_hash.update(byte_block)
        return sha256_hash.hexdigest()

    @staticmethod
    def has_macros(file_path: Path) -> bool:
        """Detecta macros VBA embebidas (p. ej. un .xlsm renombrado a .xlsx)."""
        try:
            with ZipFile(file_path, 'r') as zip_ref:
                names = zip_ref.namelist()
                return 'xl/vbaProject.bin' in names or 'word/vbaProject.bin' in names
        except Exception:
            return False

    @staticmethod
    def read_current_tag(file_path: Path) -> str:
        """Lee la etiqueta (palabras clave) que el documento ya tenga, si aplica."""
        suffix = file_path.suffix.lower()
        try:
            if suffix == '.xlsx':
                wb = load_workbook(file_path, read_only=True)
                value = wb.properties.keywords or ""
                wb.close()
                return value
            if suffix == '.docx':
                return Document(str(file_path)).core_properties.keywords or ""
            if suffix == '.pptx':
                return Presentation(str(file_path)).core_properties.keywords or ""
            if suffix == '.pdf':
                reader = PdfReader(str(file_path))
                meta = reader.metadata or {}
                return meta.get('/Keywords', '') or ''
        except Exception:
            return ""
        return ""


# ==================== MOTOR DE ETIQUETADO (sin UI, probable de forma aislada) ====================
class TaggingEngine:
    """Contiene toda la lógica de recorrido, respaldo y escritura de metadatos.
    No depende de Tkinter, así que se puede probar sin abrir ninguna ventana."""

    def __init__(self, source_root: Path, output_root: Path):
        self.source_root = source_root
        self.output_root = output_root
        self.backups_root = self._resolve_backups_root(source_root)

    @staticmethod
    def _resolve_backups_root(source_root: Path) -> Path:
        """Los respaldos NUNCA se guardan dentro de la carpeta origen: si esa
        carpeta está sincronizada con OneDrive/SharePoint, guardarlos ahí
        duplica el consumo de espacio en la nube y deja copias sin
        clasificar (la versión previa a etiquetar) flotando en una carpeta
        que puede estar compartida. Se guardan en una carpeta fija y visible
        del perfil del usuario, identificadas por un hash de la ruta origen
        para no chocar entre distintas carpetas."""
        digest = hashlib.sha256(str(source_root.resolve()).lower().encode("utf-8")).hexdigest()[:12]
        folder_name = f"{source_root.name}_{digest}"
        return get_backups_base_root() / folder_name

    def prepare_backup_dir(self):
        self.backups_root.mkdir(parents=True, exist_ok=True)

    def get_backups_size(self) -> int:
        if not self.backups_root.exists():
            return 0
        return sum(f.stat().st_size for f in self.backups_root.rglob('*') if f.is_file())

    def count_backup_files(self) -> int:
        if not self.backups_root.exists():
            return 0
        return sum(1 for f in self.backups_root.rglob('*') if f.is_file())

    @staticmethod
    def _creation_time(st) -> float:
        """Fecha de creación del archivo de respaldo (no la del documento
        original: shutil.copy2 solo preserva st_mtime/atime del original,
        así que la fecha de creación siempre queda como la de la copia)."""
        return getattr(st, "st_birthtime", None) if hasattr(st, "st_birthtime") else st.st_ctime

    def cleanup_expired_backups(self, retention_days: int) -> Tuple[int, int]:
        """Borra los respaldos creados hace más de `retention_days` días."""
        if retention_days <= 0 or not self.backups_root.exists():
            return 0, 0
        cutoff = time.time() - retention_days * 86400
        deleted, freed = 0, 0
        for path in list(self.backups_root.rglob('*')):
            if path.is_file():
                try:
                    st = path.stat()
                    if self._creation_time(st) < cutoff:
                        freed += st.st_size
                        path.unlink()
                        deleted += 1
                except OSError:
                    continue
        for dirpath, _, _ in os.walk(self.backups_root, topdown=False):
            d = Path(dirpath)
            if d == self.backups_root:
                continue
            try:
                if not any(d.iterdir()):
                    d.rmdir()
            except OSError:
                pass
        return deleted, freed

    def purge_all_backups(self) -> Tuple[int, int]:
        """Borra TODOS los respaldos de esta carpeta origen. Uso manual, solo
        después de que el usuario validó los resultados del etiquetado."""
        if not self.backups_root.exists():
            return 0, 0
        count = self.count_backup_files()
        size = self.get_backups_size()
        shutil.rmtree(self.backups_root, ignore_errors=True)
        return count, size

    def iter_candidate_files(self):
        """Recorre la carpeta de forma recursiva evitando re-escanear lo que
        el propio script genera (respaldos, logs) o la carpeta de reportes."""
        excluded_roots = set()
        try:
            resolved_output = self.output_root.resolve()
            resolved_source = self.source_root.resolve()
            if str(resolved_output).lower().startswith(str(resolved_source).lower()):
                excluded_roots.add(resolved_output)
        except Exception:
            pass

        for dirpath, dirnames, filenames in os.walk(self.source_root):
            current_dir = Path(dirpath).resolve()
            dirnames[:] = [
                d for d in dirnames
                if d.lower() not in EXCLUDED_DIR_NAMES
                and (current_dir / d) not in excluded_roots
            ]

            for name in filenames:
                # Archivos de bloqueo temporal de Office (se crean al editar un doc)
                if name.startswith('~$'):
                    continue
                # Restos de un mecanismo de respaldo anterior (copias con
                # prefijo ".backup_" junto al original) y los reportes que
                # este mismo script genera, por si la salida cae dentro del
                # origen: ninguno de los dos debe tratarse como documento.
                if name.startswith('.backup_') or name.startswith('Reporte_Ficohsa_'):
                    continue
                file_path = Path(dirpath) / name
                if file_path.suffix.lower() in SecurityValidator.ALLOWED_EXTENSIONS:
                    yield file_path

    def count_candidate_files(self) -> int:
        return sum(1 for _ in self.iter_candidate_files())

    def ensure_backup(self, file_path: Path):
        """Crea un respaldo único del archivo la primera vez que se modifica.
        Las corridas siguientes no vuelven a copiarlo: esto es lo que evita
        la duplicación, aunque proceses la misma carpeta de 1000 documentos
        una y otra vez, o le agregues archivos nuevos después."""
        rel = file_path.relative_to(self.source_root)
        backup_path = self.backups_root / rel
        if backup_path.exists():
            return
        backup_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(file_path, backup_path)

    # ---- Escritura de metadatos por tipo de documento ----
    @staticmethod
    def _write_xlsx(path: Path, tag_value: str):
        wb = load_workbook(path)
        wb.properties.keywords = tag_value
        wb.save(path)

    @staticmethod
    def _write_docx(path: Path, tag_value: str):
        doc = Document(str(path))
        doc.core_properties.keywords = tag_value
        doc.save(str(path))

    @staticmethod
    def _write_pptx(path: Path, tag_value: str):
        prs = Presentation(str(path))
        prs.core_properties.keywords = tag_value
        prs.save(str(path))

    @staticmethod
    def _write_pdf(path: Path, tag_value: str):
        reader = PdfReader(str(path))
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        writer.add_metadata({'/Keywords': tag_value})
        with open(path, 'wb') as output:
            writer.write(output)

    _HANDLERS = {
        '.xlsx': _write_xlsx.__func__,
        '.docx': _write_docx.__func__,
        '.pdf': _write_pdf.__func__,
        '.pptx': _write_pptx.__func__,
    }

    def apply_tag(self, file_path: Path, tag_value: str, create_backup: bool = True):
        handler = self._HANDLERS.get(file_path.suffix.lower())
        if handler is None:
            raise ValueError(f"Tipo no soportado: {file_path.suffix}")

        if create_backup:
            self.ensure_backup(file_path)

        with tempfile.TemporaryDirectory() as tmp_dir:
            temp_path = Path(tmp_dir) / file_path.name
            shutil.copy2(file_path, temp_path)
            handler(temp_path, tag_value)
            shutil.move(str(temp_path), str(file_path))


def format_duration(seconds: float) -> str:
    seconds = max(0, int(seconds))
    minutes, secs = divmod(seconds, 60)
    return f"{minutes:02d}:{secs:02d}"


# ==================== INTERFAZ ====================
class DocumentTaggerAppSecure:
    def __init__(self, root):
        self.root = root
        self.root.title("Etiquetador Ficohsa - Sistema Seguro")
        self.root.geometry("900x790")
        self.root.resizable(False, False)

        self.log_dir = Path.home() / ".ficohsa_tagger"
        self.security_log = SecurityLogger(self.log_dir)
        self.settings_path = self.log_dir / SETTINGS_FILE
        saved = self._load_settings()

        # Variables
        self.source_folder = ctk.StringVar(value=saved.get("source_folder", ""))
        self.output_folder = ctk.StringVar(value=saved.get("output_folder", ""))
        self.selected_tag = ctk.StringVar()
        self.simulation_mode = ctk.BooleanVar(value=saved.get("simulation_mode", True))
        self.attach_confidential = ctk.BooleanVar(value=saved.get("attach_confidential", False))
        self.compute_hash = ctk.BooleanVar(value=saved.get("compute_hash", True))
        self.backup_retention_days = ctk.StringVar(value=str(saved.get("backup_retention_days", DEFAULT_BACKUP_RETENTION_DAYS)))
        self.skip_backups = ctk.BooleanVar(value=saved.get("skip_backups", False))
        self.is_processing = False
        self.results = {"success": [], "error": [], "skipped": []}
        self.engine: Optional[TaggingEngine] = None

        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")

        self.tag_options = [
            "Confidencial",
            "Uso interno",
            "Uso restringido",
            "Auditoría interna",
            "Público",
            "Clasificado",
            "Para revisión",
            "Secreto"
        ]
        remembered_tag = saved.get("selected_tag")
        if remembered_tag and remembered_tag not in self.tag_options:
            self.tag_options.append(remembered_tag)
        self.selected_tag.set(remembered_tag if remembered_tag in self.tag_options else self.tag_options[0])

        self.setup_ui()
        self.security_log.info("Aplicación iniciada")

    # ==================== PERSISTENCIA DE PREFERENCIAS ====================
    def _load_settings(self) -> dict:
        if not self.settings_path.exists():
            return {}
        try:
            return json.loads(self.settings_path.read_text(encoding='utf-8'))
        except Exception:
            return {}

    def _save_settings(self):
        data = {
            "source_folder": self.source_folder.get(),
            "output_folder": self.output_folder.get(),
            "selected_tag": self.selected_tag.get(),
            "simulation_mode": self.simulation_mode.get(),
            "attach_confidential": self.attach_confidential.get(),
            "compute_hash": self.compute_hash.get(),
            "backup_retention_days": self.backup_retention_days.get(),
            "skip_backups": self.skip_backups.get(),
        }
        try:
            self.settings_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
        except Exception:
            pass

    def setup_ui(self):
        main_frame = ctk.CTkScrollableFrame(self.root, fg_color="#1a1a1a")
        main_frame.pack(fill="both", expand=True, padx=20, pady=20)

        title = ctk.CTkLabel(
            main_frame,
            text="🔐 Etiquetador Seguro Ficohsa",
            font=("Segoe UI", 18, "bold"),
            text_color="#00d4ff"
        )
        title.pack(pady=15)

        subtitle = ctk.CTkLabel(
            main_frame,
            text="Aplica una etiqueta de clasificación a los metadatos del documento",
            font=("Segoe UI", 10),
            text_color="#888888"
        )
        subtitle.pack(pady=(0, 20))

        # ========== CARPETAS ==========
        folder_frame = ctk.CTkFrame(main_frame, fg_color="#2a2a2a", corner_radius=10)
        folder_frame.pack(fill="x", pady=10)

        ctk.CTkLabel(folder_frame, text="📁 Carpeta a etiquetar:", font=("Segoe UI", 12, "bold")).pack(anchor="w", padx=15, pady=(15, 5))

        source_input_frame = ctk.CTkFrame(folder_frame, fg_color="#1a1a1a")
        source_input_frame.pack(fill="x", padx=15, pady=5)

        ctk.CTkEntry(source_input_frame, textvariable=self.source_folder, placeholder_text="Ruta de carpeta...").pack(side="left", fill="x", expand=True, padx=(0, 10))
        ctk.CTkButton(source_input_frame, text="Examinar", width=90, command=self.select_source_folder).pack(side="left", padx=(0, 5))
        ctk.CTkButton(source_input_frame, text="🔍 Analizar", width=100, command=self.analyze_source_folder).pack(side="left")

        self.analysis_label = ctk.CTkLabel(
            folder_frame, text="Sin analizar todavía — pulsa \"Analizar\" para ver cuántos documentos hay.",
            font=("Segoe UI", 9), text_color="#888888"
        )
        self.analysis_label.pack(anchor="w", padx=15, pady=(2, 10))

        ctk.CTkLabel(folder_frame, text="💾 Guardar reporte en:", font=("Segoe UI", 12, "bold")).pack(anchor="w", padx=15, pady=(5, 5))

        output_input_frame = ctk.CTkFrame(folder_frame, fg_color="#1a1a1a")
        output_input_frame.pack(fill="x", padx=15, pady=5)

        ctk.CTkEntry(output_input_frame, textvariable=self.output_folder, placeholder_text="Ruta de salida...").pack(side="left", fill="x", expand=True, padx=(0, 10))
        ctk.CTkButton(output_input_frame, text="Examinar", width=100, command=self.select_output_folder).pack(side="left")

        folder_frame.pack_configure(pady=(10, 5))

        # ========== ETIQUETAS ==========
        tag_frame = ctk.CTkFrame(main_frame, fg_color="#2a2a2a", corner_radius=10)
        tag_frame.pack(fill="x", pady=10)

        ctk.CTkLabel(tag_frame, text="🏷️  Seleccionar etiqueta:", font=("Segoe UI", 12, "bold")).pack(anchor="w", padx=15, pady=(15, 5))

        tag_input_frame = ctk.CTkFrame(tag_frame, fg_color="#1a1a1a")
        tag_input_frame.pack(fill="x", padx=15, pady=5)

        self.tag_combo = ctk.CTkComboBox(
            tag_input_frame,
            values=self.tag_options,
            variable=self.selected_tag,
            state="readonly"
        )
        self.tag_combo.pack(side="left", fill="x", expand=True, padx=(0, 10))

        ctk.CTkButton(tag_input_frame, text="+ Agregar", width=100, command=self.add_custom_tag).pack(side="left")

        self.confidential_check = ctk.CTkCheckBox(
            tag_frame,
            text="🔒 Adjuntar también la etiqueta \"Confidencial\" (además de la seleccionada arriba)",
            variable=self.attach_confidential,
            font=("Segoe UI", 10)
        )
        self.confidential_check.pack(anchor="w", padx=15, pady=(5, 15))

        tag_frame.pack_configure(pady=(5, 10))

        # ========== CONTROLES ==========
        controls_frame = ctk.CTkFrame(main_frame, fg_color="#2a2a2a", corner_radius=10)
        controls_frame.pack(fill="x", pady=10)

        ctk.CTkLabel(controls_frame, text="⚙️  Controles:", font=("Segoe UI", 12, "bold")).pack(anchor="w", padx=15, pady=(15, 10))

        self.simulation_check = ctk.CTkCheckBox(
            controls_frame,
            text="🔍 Modo Simulación (vista previa: no modifica archivos)",
            variable=self.simulation_mode,
            font=("Segoe UI", 10)
        )
        self.simulation_check.pack(anchor="w", padx=15, pady=5)

        self.hash_check = ctk.CTkCheckBox(
            controls_frame,
            text="🧮 Calcular hash SHA256 de auditoría (desactívalo en lotes grandes para ir más rápido)",
            variable=self.compute_hash,
            font=("Segoe UI", 10)
        )
        self.hash_check.pack(anchor="w", padx=15, pady=5)

        self.skip_backups_check = ctk.CTkCheckBox(
            controls_frame,
            text="🚫 No generar respaldos (más rápido, pero sin red de seguridad para deshacer cambios)",
            variable=self.skip_backups,
            font=("Segoe UI", 10),
            text_color="#ff8888"
        )
        self.skip_backups_check.pack(anchor="w", padx=15, pady=5)

        retention_frame = ctk.CTkFrame(controls_frame, fg_color="transparent")
        retention_frame.pack(fill="x", padx=15, pady=5)
        ctk.CTkLabel(retention_frame, text="🗓️ Retención de respaldos (días):", font=("Segoe UI", 10)).pack(side="left")
        ctk.CTkEntry(retention_frame, textvariable=self.backup_retention_days, width=60).pack(side="left", padx=(10, 0))

        ctk.CTkLabel(
            controls_frame,
            text="Los archivos ya etiquetados se omiten automáticamente en cada corrida para evitar\n"
                 "duplicados. Los respaldos se guardan en la carpeta visible \"FicohsaBackups\" de tu\n"
                 "perfil de usuario (fuera de OneDrive) y se limpian solos al vencer la retención\n"
                 "configurada arriba.",
            font=("Segoe UI", 9),
            text_color="#888888",
            justify="left"
        ).pack(anchor="w", padx=15, pady=(0, 15))

        controls_frame.pack_configure(pady=(10, 5))

        # ========== PROGRESO ==========
        progress_frame = ctk.CTkFrame(main_frame, fg_color="transparent")
        progress_frame.pack(fill="x", pady=15)

        self.progress = ctk.CTkProgressBar(progress_frame, height=8)
        self.progress.pack(fill="x")
        self.progress.set(0)

        self.status_label = ctk.CTkLabel(
            progress_frame,
            text="🟢 Listo",
            font=("Segoe UI", 10),
            text_color="#00ff00"
        )
        self.status_label.pack(anchor="w", pady=(5, 0))

        # ========== BOTONES ==========
        button_frame = ctk.CTkFrame(main_frame, fg_color="transparent")
        button_frame.pack(fill="x", pady=15)

        self.start_button = ctk.CTkButton(
            button_frame,
            text="▶️  Iniciar Etiquetado",
            command=self.start_tagging,
            font=("Segoe UI", 12, "bold"),
            fg_color="#00d4ff",
            hover_color="#0099cc"
        )
        self.start_button.pack(side="left", padx=5)

        self.stop_button = ctk.CTkButton(
            button_frame,
            text="⏹️  Cancelar",
            command=self.stop_tagging,
            state="disabled",
            font=("Segoe UI", 12, "bold"),
            fg_color="#ff4444",
            hover_color="#cc0000"
        )
        self.stop_button.pack(side="left", padx=5)

        ctk.CTkButton(
            button_frame,
            text="📜 Ver Log",
            command=self.open_log_file,
            font=("Segoe UI", 11),
            fg_color="#555555",
            hover_color="#777777"
        ).pack(side="left", padx=5)

        self.backups_button = ctk.CTkButton(
            button_frame,
            text="🗄️ Ver Respaldos",
            command=self.open_backups_folder,
            font=("Segoe UI", 11),
            fg_color="#555555",
            hover_color="#777777"
        )
        self.backups_button.pack(side="left", padx=5)

        self.purge_button = ctk.CTkButton(
            button_frame,
            text="🧹 Purgar Respaldos",
            command=self.purge_backups,
            font=("Segoe UI", 11),
            fg_color="#8a4b00",
            hover_color="#b35f00"
        )
        self.purge_button.pack(side="left", padx=5)

        self.report_button = ctk.CTkButton(
            button_frame,
            text="📂 Ver Reporte",
            command=self.open_output_folder,
            font=("Segoe UI", 11),
            fg_color="#555555",
            hover_color="#777777"
        )
        self.report_button.pack(side="left", padx=5)

        # ========== LOG ==========
        ctk.CTkLabel(main_frame, text="📋 Actividad:", font=("Segoe UI", 12, "bold")).pack(anchor="w", pady=(10, 5))

        self.log_text = ctk.CTkTextbox(main_frame, height=200, fg_color="#0a0a0a", text_color="#00ff00", font=("Courier New", 9))
        self.log_text.pack(fill="both", expand=True, pady=5)

    def log_ui(self, message: str, level: str = "INFO"):
        timestamp = datetime.now().strftime('%H:%M:%S')

        color_map = {
            "INFO": "#00ff00",
            "ERROR": "#ff4444",
            "WARNING": "#ffaa00",
            "SUCCESS": "#00ff00",
            "DEBUG": "#888888"
        }

        self.log_text.insert("end", f"[{timestamp}] {message}\n", color_map.get(level, "#ffffff"))
        self.log_text.see("end")

        log_method = {
            "INFO": self.security_log.info,
            "ERROR": self.security_log.error,
            "WARNING": self.security_log.warning,
            "DEBUG": self.security_log.debug,
            "SUCCESS": self.security_log.info
        }.get(level, self.security_log.info)

        log_method(message)

    def select_source_folder(self):
        folder = filedialog.askdirectory(title="Selecciona la carpeta a etiquetar")
        if folder:
            self.source_folder.set(folder)
            self.analysis_label.configure(text="Sin analizar todavía — pulsa \"Analizar\" para ver cuántos documentos hay.")
            self.log_ui(f"Carpeta origen: {folder}", "INFO")

    def select_output_folder(self):
        folder = filedialog.askdirectory(title="Selecciona donde guardar el reporte")
        if folder:
            self.output_folder.set(folder)
            self.log_ui(f"Carpeta salida: {folder}", "INFO")

    def analyze_source_folder(self):
        """Cuenta los documentos compatibles sin modificar nada — útil antes
        de lanzar un lote grande (p. ej. 1000 documentos) para saber qué esperar."""
        source = self.source_folder.get()
        if not source or not Path(source).exists():
            messagebox.showerror("Error", "Selecciona primero una carpeta válida")
            return

        self.analysis_label.configure(text="🔎 Analizando...")

        def run():
            engine = TaggingEngine(Path(source), Path(self.output_folder.get() or source))
            counts = Counter(f.suffix.lower() for f in engine.iter_candidate_files())
            total = sum(counts.values())
            detail = ", ".join(f"{n} {ext}" for ext, n in sorted(counts.items())) or "ninguno"
            text = f"Encontrados: {total} documentos compatibles ({detail})"
            self.root.after(0, lambda: self.analysis_label.configure(text=text))
            self.root.after(0, lambda: self.log_ui(f"🔍 Análisis: {text}", "INFO"))

        threading.Thread(target=run, daemon=True).start()

    def add_custom_tag(self):
        dialog = ctk.CTkInputDialog(text="Nueva etiqueta:", title="Agregar Etiqueta")
        new_tag = dialog.get_input()

        if new_tag:
            new_tag = new_tag.strip().capitalize()
            if new_tag and new_tag not in self.tag_options:
                self.tag_options.append(new_tag)
                self.tag_combo.configure(values=self.tag_options)
                self.selected_tag.set(new_tag)
                self.log_ui(f"✓ Etiqueta agregada: {new_tag}", "SUCCESS")
                self.security_log.info(f"Nueva etiqueta creada: {new_tag}")
            else:
                messagebox.showwarning("Advertencia", "Etiqueta duplicada o vacía")

    def start_tagging(self):
        if not self.source_folder.get() or not self.output_folder.get():
            messagebox.showerror("Error", "Selecciona carpetas origen y salida")
            return

        if not Path(self.source_folder.get()).exists():
            messagebox.showerror("Error", "La carpeta a etiquetar no existe")
            return

        if not self.simulation_mode.get():
            if self.skip_backups.get():
                backup_text = (
                    "⚠️ Tienes activada la opción \"No generar respaldos\": si algo sale mal "
                    "durante el etiquetado, NO habrá copia previa para restaurar el archivo.\n\n"
                )
            else:
                backup_text = (
                    "Se creará un respaldo único por cada archivo que se modifique por primera vez, "
                    "guardado en la carpeta visible \"FicohsaBackups\" de tu perfil de usuario "
                    "(fuera de OneDrive, no dentro de esta carpeta). Los respaldos "
                    f"con más de {self.backup_retention_days.get()} días se eliminan automáticamente al "
                    "finalizar cada corrida.\n\n"
                )
            confirmed = messagebox.askyesno(
                "Confirmar cambios reales",
                "Vas a modificar los metadatos de los documentos reales en:\n"
                f"{self.source_folder.get()}\n\n"
                f"Etiqueta: {self.selected_tag.get()}" + (" + Confidencial" if self.attach_confidential.get() else "") + "\n\n"
                + backup_text +
                "Los archivos que ya tengan esta etiqueta se omiten y no se vuelven a tocar.\n\n"
                "¿Deseas continuar?"
            )
            if not confirmed:
                return

        self._save_settings()

        self.is_processing = True
        self.start_button.configure(state="disabled")
        self.stop_button.configure(state="normal")
        self.status_label.configure(text="🔵 Procesando...", text_color="#ffaa00")
        self.results = {"success": [], "error": [], "skipped": []}

        self.security_log.info(f"Iniciando etiquetado - Modo simulación: {self.simulation_mode.get()}")

        thread = threading.Thread(target=self.tag_documents)
        thread.daemon = True
        thread.start()

    def stop_tagging(self):
        self.is_processing = False
        self.log_ui("❌ Cancelado por usuario", "WARNING")

    def tag_documents(self):
        start_time = time.time()
        try:
            source_path = Path(self.source_folder.get())
            output_path = Path(self.output_folder.get())
            tag_value = self.selected_tag.get()
            sim_mode = self.simulation_mode.get()
            attach_confidential = self.attach_confidential.get()
            use_hash = self.compute_hash.get()
            make_backups = not self.skip_backups.get()

            self.engine = TaggingEngine(source_path, output_path)
            if not sim_mode and make_backups:
                self.engine.prepare_backup_dir()

            files = list(self.engine.iter_candidate_files())

            self.log_ui(f"📊 Encontrados {len(files)} documentos compatibles (.xlsx, .docx, .pdf, .pptx)", "INFO")
            self.log_ui(f"🏷️  Etiqueta: {tag_value}" + (" + Confidencial" if attach_confidential else ""), "INFO")

            if not sim_mode and not make_backups:
                self.log_ui("🚫 Respaldos desactivados: los archivos se modificarán sin copia previa.", "WARNING")

            old_backups = source_path / BACKUPS_DIRNAME
            if old_backups.exists():
                self.log_ui(
                    "⚠️ Se detectó una carpeta \".ficohsa_backups\" antigua dentro del origen (versión "
                    "anterior del script). Los respaldos nuevos ahora se guardan localmente fuera de "
                    "OneDrive; revisa y elimina esa carpeta antigua manualmente si ya no la necesitas.",
                    "WARNING"
                )

            if sim_mode:
                self.log_ui("⚠️  MODO SIMULACIÓN ACTIVO - No se modificarán archivos", "WARNING")

            if not files:
                self.log_ui("⚠️  No se encontraron documentos", "WARNING")
                self.progress.set(1.0)
                self.status_label.configure(text="✅ Completado", text_color="#00ff00")
                return

            total = len(files)
            for idx, file_path in enumerate(files):
                if not self.is_processing:
                    break

                progress = (idx + 1) / total
                self.progress.set(progress)

                elapsed = time.time() - start_time
                rate = (idx + 1) / elapsed if elapsed > 0 else 0
                eta = (total - idx - 1) / rate if rate > 0 else 0
                self.status_label.configure(
                    text=f"🔵 {idx + 1}/{total} | ✓{len(self.results['success'])} "
                         f"✗{len(self.results['error'])} ⚠{len(self.results['skipped'])} | "
                         f"transcurrido {format_duration(elapsed)} · restante ~{format_duration(eta)}",
                    text_color="#ffaa00"
                )

                try:
                    self.log_ui(f"[{idx+1}/{total}] Procesando: {file_path.name}", "INFO")

                    is_valid, validation_msg = SecurityValidator.validate_file(file_path)
                    if not is_valid:
                        self._skip(file_path, validation_msg)
                        continue

                    if SecurityValidator.is_cloud_placeholder(file_path):
                        self._skip(file_path, "Archivo solo en la nube (OneDrive) - no descargado localmente")
                        continue

                    if SecurityValidator.has_macros(file_path):
                        self._skip(file_path, "⚠️ Archivo contiene macros VBA - omitido por seguridad")
                        continue

                    existing_tag = SecurityValidator.read_current_tag(file_path)
                    existing_tags = {t.strip() for t in existing_tag.split(",") if t.strip()}
                    desired_tags = {tag_value.strip()}
                    if attach_confidential:
                        desired_tags.add("Confidencial")

                    if desired_tags.issubset(existing_tags):
                        self._skip(file_path, "Ya etiquetado con esta clasificación")
                        continue

                    final_tag_value = ", ".join(sorted(existing_tags | desired_tags))

                    hash_before = SecurityValidator.calculate_hash(file_path) if use_hash else ""

                    if not sim_mode:
                        self.engine.apply_tag(file_path, final_tag_value, create_backup=make_backups)
                        hash_after = SecurityValidator.calculate_hash(file_path) if use_hash else ""
                    else:
                        hash_after = hash_before

                    self.results["success"].append({
                        "archivo": file_path.name,
                        "ruta": str(file_path),
                        "tipo": file_path.suffix,
                        "etiqueta": final_tag_value,
                        "hash_antes": hash_before,
                        "hash_después": hash_after,
                        "simulación": sim_mode,
                        "timestamp": datetime.now().isoformat()
                    })

                    self.log_ui("  ✓ Completado", "SUCCESS")

                except Exception as e:
                    self.log_ui(f"  ✗ Error: {str(e)}", "ERROR")
                    self.results["error"].append({
                        "archivo": file_path.name,
                        "ruta": str(file_path),
                        "error": str(e),
                        "timestamp": datetime.now().isoformat()
                    })

            self.log_ui("📊 Generando reporte...", "INFO")
            self._generate_secure_report()

            if not sim_mode:
                self._audit_and_cleanup_backups()

            self.progress.set(1.0)
            self.status_label.configure(text="✅ Completado", text_color="#00ff00")

            skip_breakdown = Counter(r["razón"] for r in self.results["skipped"])
            for reason, count in skip_breakdown.most_common():
                self.log_ui(f"  ⚠️  {count}x: {reason}", "INFO")

            self.log_ui("✅ Proceso finalizado", "SUCCESS")

            breakdown_text = "\n".join(f"  • {count}x {reason}" for reason, count in skip_breakdown.most_common(5))
            self.root.after(0, lambda: messagebox.showinfo(
                "Éxito",
                f"✅ Etiquetado completado en {format_duration(time.time() - start_time)}\n\n"
                f"Exitosos: {len(self.results['success'])}\n"
                f"Errores: {len(self.results['error'])}\n"
                f"Omitidos: {len(self.results['skipped'])}\n"
                + (f"\nMotivos de omisión:\n{breakdown_text}" if breakdown_text else "")
            ))

        except Exception as e:
            self.log_ui(f"❌ ERROR CRÍTICO: {str(e)}", "ERROR")
            self.status_label.configure(text="🔴 Error", text_color="#ff4444")
            self.root.after(0, lambda: messagebox.showerror("Error Crítico", str(e)))

        finally:
            self.start_button.configure(state="normal")
            self.stop_button.configure(state="disabled")

    def _skip(self, file_path: Path, reason: str):
        self.log_ui(f"  ⚠️  Omitido: {reason}", "WARNING")
        self.results["skipped"].append({"archivo": file_path.name, "razón": reason})

    def _audit_and_cleanup_backups(self):
        """Registra cuánto espacio ocupan los respaldos (para poder auditarlo)
        y aplica la política de retención configurada en la UI."""
        if self.engine is None:
            return

        count = self.engine.count_backup_files()
        size = self.engine.get_backups_size()
        self.log_ui(f"🗄️ Respaldos actuales: {count} archivos, {format_size(size)} en {self.engine.backups_root}", "INFO")
        self.security_log.info(
            f"AUDITORÍA_RESPALDOS | archivos={count} | tamaño_bytes={size} | ruta={self.engine.backups_root}"
        )

        try:
            retention_days = int(self.backup_retention_days.get())
        except ValueError:
            retention_days = DEFAULT_BACKUP_RETENTION_DAYS

        deleted, freed = self.engine.cleanup_expired_backups(retention_days)
        if deleted:
            self.log_ui(
                f"🧹 Limpieza automática: {deleted} respaldos con más de {retention_days} días eliminados "
                f"({format_size(freed)} liberados)", "INFO"
            )
            self.security_log.info(
                f"LIMPIEZA_AUTOMATICA_RESPALDOS | eliminados={deleted} | bytes_liberados={freed} | "
                f"retención_días={retention_days}"
            )

    def _generate_secure_report(self):
        try:
            output_path = Path(self.output_folder.get())
            report_file = output_path / f"Reporte_Ficohsa_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

            wb = Workbook()

            ws_summary = wb.active
            ws_summary.title = "Resumen"

            ws_summary['A1'] = "REPORTE DE ETIQUETADO - FICOHSA"
            ws_summary['A3'] = f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            ws_summary['A4'] = f"Etiqueta: {self.selected_tag.get()}"
            ws_summary['A5'] = f"Modo Simulación: {'SÍ' if self.simulation_mode.get() else 'NO'}"
            ws_summary['A6'] = f"Total procesados: {len(self.results['success']) + len(self.results['error']) + len(self.results['skipped'])}"
            ws_summary['A7'] = f"Exitosos: {len(self.results['success'])}"
            ws_summary['A8'] = f"Errores: {len(self.results['error'])}"
            ws_summary['A9'] = f"Omitidos: {len(self.results['skipped'])}"

            if self.results["success"]:
                ws_success = wb.create_sheet("Exitosos")
                headers = ["Archivo", "Ruta", "Tipo", "Etiqueta", "Hash Antes", "Hash Después", "Simulación", "Timestamp"]
                ws_success.append(headers)

                for result in self.results["success"]:
                    ws_success.append([
                        result["archivo"],
                        result["ruta"],
                        result["tipo"],
                        result["etiqueta"],
                        (result["hash_antes"][:16] + "...") if result["hash_antes"] else "(sin calcular)",
                        (result["hash_después"][:16] + "...") if result["hash_después"] else "(sin calcular)",
                        "SÍ" if result["simulación"] else "NO",
                        result["timestamp"]
                    ])

            if self.results["error"]:
                ws_error = wb.create_sheet("Errores")
                headers = ["Archivo", "Ruta", "Error", "Timestamp"]
                ws_error.append(headers)

                for result in self.results["error"]:
                    ws_error.append([
                        result["archivo"],
                        result["ruta"],
                        result["error"],
                        result["timestamp"]
                    ])

            if self.results["skipped"]:
                ws_skipped = wb.create_sheet("Omitidos")
                headers = ["Archivo", "Razón"]
                ws_skipped.append(headers)

                for result in self.results["skipped"]:
                    ws_skipped.append([
                        result["archivo"],
                        result["razón"]
                    ])

            wb.save(str(report_file))
            self.log_ui(f"💾 Reporte: {report_file.name}", "SUCCESS")

        except Exception as e:
            self.log_ui(f"Error generando reporte: {str(e)}", "ERROR")

    def open_log_file(self):
        log_file = self.security_log.log_file
        if log_file.exists():
            os.startfile(str(log_file)) if sys.platform == 'win32' else os.system(f"open '{log_file}'")

    def open_backups_folder(self):
        source = self.source_folder.get()
        if not source:
            messagebox.showinfo("Respaldos", "Primero selecciona una carpeta a etiquetar")
            return
        engine = TaggingEngine(Path(source), Path(self.output_folder.get() or source))
        backups_path = engine.backups_root
        if not backups_path.exists():
            messagebox.showinfo("Respaldos", "Todavía no hay respaldos para esta carpeta (aún no se ha modificado ningún archivo).")
            return
        os.startfile(str(backups_path)) if sys.platform == 'win32' else os.system(f"open '{backups_path}'")

    def purge_backups(self):
        """Borra manualmente todos los respaldos de la carpeta origen actual.
        Pensado para usarse después de validar el reporte y revisar algunos
        documentos al azar — no hay forma de deshacer esta acción."""
        source = self.source_folder.get()
        if not source:
            messagebox.showinfo("Respaldos", "Primero selecciona una carpeta a etiquetar")
            return

        engine = TaggingEngine(Path(source), Path(self.output_folder.get() or source))
        if not engine.backups_root.exists():
            messagebox.showinfo("Respaldos", "No hay respaldos que eliminar para esta carpeta.")
            return

        count = engine.count_backup_files()
        size = engine.get_backups_size()
        confirmed = messagebox.askyesno(
            "Confirmar eliminación de respaldos",
            f"Se eliminarán {count} archivos de respaldo ({format_size(size)}) de:\n{engine.backups_root}\n\n"
            "Hazlo solo después de validar que los documentos quedaron etiquetados correctamente "
            "(revisa el reporte y abre algunos documentos al azar).\n\n"
            "Esta acción NO se puede deshacer. ¿Continuar?"
        )
        if not confirmed:
            return

        deleted, freed = engine.purge_all_backups()
        user = os.environ.get("USERNAME", "desconocido")
        self.security_log.info(
            f"PURGA_MANUAL_RESPALDOS | usuario={user} | archivos={deleted} | bytes_liberados={freed} | "
            f"ruta={engine.backups_root}"
        )
        self.log_ui(f"🧹 Respaldos eliminados manualmente: {deleted} archivos ({format_size(freed)})", "SUCCESS")
        messagebox.showinfo("Respaldos eliminados", f"Se eliminaron {deleted} archivos ({format_size(freed)}).")

    def open_output_folder(self):
        output = self.output_folder.get()
        if not output or not Path(output).exists():
            messagebox.showinfo("Reporte", "Primero selecciona una carpeta de salida válida")
            return
        os.startfile(output) if sys.platform == 'win32' else os.system(f"open '{output}'")


# ==================== MAIN ====================
if __name__ == "__main__":
    root = ctk.CTk()
    app = DocumentTaggerAppSecure(root)
    root.mainloop()
